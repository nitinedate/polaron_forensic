"""Load Magnet AXIOM 10.2.0 grouped workbook into public.axiom_* tables."""

from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from typing import Any

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.db.sql_helpers import rollback_aborted_transaction

log = logging.getLogger("axiom_catalog_ingest")

OS_FAMILY_TO_AXIOM_PLATFORM: dict[str, str] = {
    "windows": "Windows",
    "linux": "Linux",
    "macos": "macOS",
    "android": "Android",
    "ios": "iOS",
    "chromebook": "Chromebook",
    "cloud": "Cloud",
}

# Magnet catalog IDs whose DB label differs from the AXIOM/report template name.
_CATALOG_ARTIFACT_LABELS: dict[str, str] = {
    "AX-0011": "Installed Programs (Non-Microsoft)",
}


def catalog_artifact_label(artifact_id: str, artifact_name: str | None) -> str:
    return _CATALOG_ARTIFACT_LABELS.get(artifact_id) or artifact_name or artifact_id


def catalog_display_status(*, stored_status: str | None = None) -> str:
    """Catalog badge: unavailable/uncounted items are done+0, not pending forever.

    Inventory still overwrites the count when records appear.
    """
    st = (stored_status or "").strip().lower()
    if st in {"failed", "error"}:
        return "failed"
    return "done"


def ensure_platform_axiom_catalog(db: Session, platform: str) -> dict[str, Any]:
    """Ensure axiom_artifacts has rows for this platform (iOS/Android/Windows).

    Mobile jobs previously stalled forever when the Magnet workbook was missing and
    only a tiny Windows seed existed — inventory total stayed 0 so the supervisor
    never queued Artifact inventory.
    """
    plat = (platform or "Windows").strip() or "Windows"
    rollback_aborted_transaction(db)
    existing = db.execute(
        text("SELECT count(*) c FROM public.axiom_artifacts WHERE platform = :p"),
        {"p": plat},
    ).mappings().first()
    count = int((existing or {}).get("c") or 0)

    def _platform_count() -> int:
        after = db.execute(
            text("SELECT count(*) c FROM public.axiom_artifacts WHERE platform = :p"),
            {"p": plat},
        ).mappings().first()
        return int((after or {}).get("c") or 0)

    # Prefer Magnet workbook when present (full multi-platform catalog).
    try:
        xlsx = _workbook_path(axiom_data_dir())
        if xlsx.is_file():
            # Reload when this platform is empty even if other platforms exist.
            if count <= 0:
                loaded = load_axiom_catalog_from_files(db, force=True)
                count = _platform_count()
                if count > 0:
                    return {"platform": plat, "count": count, "source": "workbook", **(loaded or {})}
    except Exception as exc:
        log.warning("workbook catalog ensure failed platform=%s: %s", plat, exc)
        rollback_aborted_transaction(db)

    if plat in ("iOS", "Android"):
        from app.services.mobile_report_catalog import (
            MOBILE_FORENSIC_ARTIFACTS,
            MOBILE_FORENSIC_ARTIFACTS_IOS,
        )

        pairs = list(MOBILE_FORENSIC_ARTIFACTS_IOS if plat == "iOS" else MOBILE_FORENSIC_ARTIFACTS)
        try:
            with db.begin_nested():
                added = _upsert_mobile_catalog_pairs(db, plat, pairs)
            return {
                "platform": plat,
                "count": _platform_count(),
                "added": added,
                "source": "mobile_report_catalog",
            }
        except Exception as exc:
            log.warning("mobile catalog ensure failed platform=%s: %s", plat, exc)
            rollback_aborted_transaction(db)
            return {"platform": plat, "count": count, "source": "none", "error": str(exc)[:200]}

    # Windows / other — report template section-B seed
    try:
        from app.services.report_catalog_sync import ensure_report_template_artifacts

        with db.begin_nested():
            ensure_report_template_artifacts(db, platform=plat)
        return {
            "platform": plat,
            "count": _platform_count(),
            "source": "report_template",
        }
    except Exception as exc:
        log.warning("report template catalog ensure failed platform=%s: %s", plat, exc)
        rollback_aborted_transaction(db)
        return {"platform": plat, "count": count, "source": "none", "error": str(exc)[:200]}


def _upsert_mobile_catalog_pairs(
    db: Session, platform: str, pairs: list[tuple[str, str]]
) -> int:
    """Insert mobile forensic artifact names for a platform (idempotent by name)."""
    import hashlib

    existing_names = {
        (r["artifact_name"] or "").strip().lower()
        for r in db.execute(
            text("SELECT artifact_name FROM public.axiom_artifacts WHERE platform = :p"),
            {"p": platform},
        ).mappings().all()
    }
    max_sort = db.execute(
        text("SELECT coalesce(max(sort_order), 0) m FROM public.axiom_artifacts WHERE platform = :p"),
        {"p": platform},
    ).mappings().first()
    sort_order = int((max_sort or {}).get("m") or 0)
    added = 0
    prefix = "IOS" if platform == "iOS" else "AND"
    for category, name in pairs:
        key = (name or "").strip().lower()
        if not key or key in existing_names:
            continue
        sort_order += 1
        digest = hashlib.sha1(f"{platform}:{category}:{name}".encode("utf-8")).hexdigest()[:8].upper()
        aid = f"MOB-{prefix}-{digest}"
        prompt = (
            f"For the {platform} forensic artifact \"{name}\" ({category}): "
            f"report the total count recovered on this evidence, a concise explanation of what was found, "
            f"and the forensic significance."
        )
        db.execute(
            text(
                """INSERT INTO public.axiom_artifacts
                   (artifact_id, platform, category, artifact_name, recovery_method,
                    prompt_question, critical, sort_order, observation_focus, metadata, updated_at)
                   VALUES (:aid, :platform, :category, :name, 'Parsing', :prompt, FALSE, :sort,
                           :obs, CAST(:meta AS jsonb), NOW())
                   ON CONFLICT (artifact_id) DO UPDATE SET
                     platform = EXCLUDED.platform,
                     category = EXCLUDED.category,
                     artifact_name = EXCLUDED.artifact_name,
                     prompt_question = EXCLUDED.prompt_question,
                     observation_focus = EXCLUDED.observation_focus,
                     updated_at = NOW()"""
            ),
            {
                "aid": aid,
                "platform": platform,
                "category": category,
                "name": name,
                "prompt": prompt,
                "sort": sort_order,
                "obs": f"{platform} {name}",
                "meta": json.dumps(
                    {"source": "mobile_report_catalog", "platform": platform, "category": category}
                ),
            },
        )
        existing_names.add(key)
        added += 1
    if added:
        try:
            db.flush()
        except Exception:
            pass
    return added


def axiom_data_dir() -> Path:
    here = Path(__file__).resolve()
    candidates: list[Path] = []
    try:
        from app.config import get_settings

        candidates.append(Path(get_settings().data_root) / "axiom")
    except Exception:
        pass
    for n in (2, 3, 4):
        if n < len(here.parents):
            candidates.append(here.parents[n] / "data" / "axiom")
    candidates.extend((Path("/app/data/axiom"), Path("/data/axiom")))
    seen: set[str] = set()
    for candidate in candidates:
        key = str(candidate)
        if key in seen:
            continue
        seen.add(key)
        if candidate.is_dir():
            return candidate
    return candidates[0] if candidates else Path("/app/data/axiom")


def _workbook_path(data_dir: Path) -> Path:
    for name in (
        "grouped_workbook.xlsx",
        "Magnet_AXIOM_10.2.0_Grouped_Artifact_Objective_Procedure_Workbook.xlsx",
        "catalog.xlsx",
    ):
        path = data_dir / name
        if path.is_file():
            return path
    return data_dir / "grouped_workbook.xlsx"


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text_val = str(value).strip()
    return text_val or None


def _build_artifact_prompt(
    *,
    platform: str,
    category: str,
    artifact_name: str,
    observation_focus: str | None,
    objective: dict[str, Any] | None,
    procedure: dict[str, Any] | None,
    recovery_method: str | None,
) -> str:
    obj_title = (objective or {}).get("title") or "Forensic examination"
    obj_statement = (objective or {}).get("statement") or ""
    obs_fields = (
        (procedure or {}).get("expected_output_fields")
        or (objective or {}).get("required_observation_fields")
        or observation_focus
        or ""
    )
    corroboration = (procedure or {}).get("mandatory_corroboration") or (objective or {}).get("minimum_corroboration") or ""
    return (
        f"For the {platform} Magnet AXIOM artifact \"{artifact_name}\" ({category}): "
        f"report the total count recovered on this evidence, a concise explanation of what was found, "
        f"and the forensic significance. Include sample source paths or databases, key timestamps, "
        f"and users/accounts where available. "
        f"Observation fields: {obs_fields}. "
        f"Corroboration: {corroboration}. "
        f"Objective — {obj_title}: {obj_statement} "
        f"Recovery method: {recovery_method or 'Parsing'}."
    ).strip()


def _build_objective_prompt(obj: dict[str, Any], procedure: dict[str, Any] | None) -> str:
    obs = obj.get("required_observation_fields") or ""
    proc_title = (procedure or {}).get("title") or ""
    proc_fields = (procedure or {}).get("expected_output_fields") or obs
    return (
        f"Objective {obj.get('objective_id')} — {obj.get('title')}: {obj.get('statement') or ''} "
        f"Required observations: {obs}. "
        f"Minimum corroboration: {obj.get('minimum_corroboration') or 'N/A'}. "
        f"Linked procedure {obj.get('procedure_id') or (procedure or {}).get('procedure_id') or ''}"
        f"{(' — ' + proc_title) if proc_title else ''}. "
        f"Expected output fields: {proc_fields}. "
        f"Report counts, observed facts, and forensic significance grounded in indexed evidence."
    ).strip()


def _build_procedure_prompt(proc: dict[str, Any], objective: dict[str, Any] | None) -> str:
    obj_title = (objective or {}).get("title") or ""
    return (
        f"Procedure {proc.get('procedure_id')} for objective {(objective or {}).get('objective_id') or proc.get('objective_id')}"
        f"{(' — ' + obj_title) if obj_title else ''}: "
        f"{proc.get('title') or ''}. "
        f"Expected observations: {proc.get('expected_output_fields') or 'N/A'}. "
        f"Mandatory corroboration: {proc.get('mandatory_corroboration') or 'N/A'}. "
        f"Follow the examination steps and report counts, observed facts, corroboration, and limitations "
        f"using only indexed evidence for this job."
    ).strip()


def _sheet_rows(xlsx_path: Path, sheet: str) -> tuple[list[str], list[tuple[Any, ...]]]:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[sheet]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()
    return headers, rows


def _load_objectives(xlsx_path: Path) -> dict[str, dict[str, Any]]:
    headers, rows = _sheet_rows(xlsx_path, "Objectives")
    idx = {h: i for i, h in enumerate(headers) if h}
    out: dict[str, dict[str, Any]] = {}
    for row in rows:
        if not row or not row[idx.get("Objective ID", 0)]:
            continue
        oid = str(row[idx["Objective ID"]]).strip()
        out[oid] = {
            "objective_id": oid,
            "procedure_id": _clean(row[idx.get("Procedure ID", 1)]),
            "domain": _clean(row[idx.get("Domain", 2)]),
            "title": _clean(row[idx.get("Objective Title", 3)]),
            "statement": _clean(row[idx.get("Objective Statement", 4)]),
            "primary_artifact_families": _clean(row[idx.get("Primary Artifact Families", 5)]),
            "required_observation_fields": _clean(row[idx.get("Required Observation Fields", 6)]),
            "minimum_corroboration": _clean(row[idx.get("Minimum Corroboration", 7)]),
            "limitations": _clean(row[idx.get("Limitations and Cautions", 8)]),
            "priority": _clean(row[idx.get("Priority", 9)]),
        }
    return out


def _load_procedures(xlsx_path: Path) -> dict[str, dict[str, Any]]:
    headers, rows = _sheet_rows(xlsx_path, "Procedures")
    idx = {h: i for i, h in enumerate(headers) if h}
    out: dict[str, dict[str, Any]] = {}
    for row in rows:
        if not row or not row[idx.get("Procedure ID", 0)]:
            continue
        pid = str(row[idx["Procedure ID"]]).strip()
        out[pid] = {
            "procedure_id": pid,
            "objective_id": _clean(row[idx.get("Objective ID", 1)]),
            "domain": _clean(row[idx.get("Domain", 2)]),
            "title": _clean(row[idx.get("Procedure Title", 3)]),
            "detailed_procedure": _clean(row[idx.get("Detailed Procedure", 4)]),
            "mandatory_corroboration": _clean(row[idx.get("Mandatory Corroboration", 5)]),
            "source_validation_requirement": _clean(row[idx.get("Source Validation Requirement", 6)]),
            "expected_output_fields": _clean(row[idx.get("Expected Output / Observation Fields", 7)]),
            "limitations": _clean(row[idx.get("Limitations and Cautions", 8)]),
            "minimum_confidence_rule": _clean(row[idx.get("Minimum Confidence Rule", 9)]),
        }
    return out


def _load_grouped_artifacts(xlsx_path: Path) -> list[dict[str, Any]]:
    headers, rows = _sheet_rows(xlsx_path, "All_Grouped_Artifacts")
    idx = {h: i for i, h in enumerate(headers) if h}

    def col(row: tuple, name: str, *aliases: str) -> Any:
        for key in (name, *aliases):
            if key in idx:
                return row[idx[key]]
        return None

    records: list[dict[str, Any]] = []
    for row in rows:
        artifact_id = _clean(col(row, "Catalogue ID", "artifact_id"))
        platform = _clean(col(row, "Platform", "platform"))
        category = _clean(col(row, "Category", "category"))
        artifact_name = _clean(col(row, "Native Artifact Name", "artifact_name"))
        if not artifact_id or not platform or not category or not artifact_name:
            continue
        ref_page = col(row, "Reference Page", "reference_page")
        try:
            reference_page = int(ref_page) if ref_page not in (None, "") else None
        except (TypeError, ValueError):
            reference_page = None
        records.append({
            "artifact_id": artifact_id,
            "platform": platform,
            "category": category,
            "application_or_profile": _clean(col(row, "Application / Profile", "application_or_profile")),
            "artifact_name": artifact_name,
            "recovery_method": _clean(col(row, "Recovery Method", "recovery_method")),
            "reference_page": reference_page,
            "primary_objective_id": _clean(col(row, "Primary Objective ID", "primary_objective_id")),
            "secondary_objective_ids": _clean(col(row, "Secondary Objective IDs", "secondary_objective_ids")),
            "procedure_id": _clean(col(row, "Procedure ID", "procedure_id")),
            "observation_focus": _clean(col(row, "Observation Focus", "observation_focus")),
            "outline_path": _clean(col(row, "Outline Path", "outline_path")),
            "source_version": _clean(col(row, "Source Version", "source_version")),
            "source_published": _clean(col(row, "Source Published", "source_published")),
            "source_url": _clean(col(row, "Source URL", "source_url")),
            "mapping_note": _clean(col(row, "Mapping Note", "mapping_note")),
        })
    return records


def _upsert_objectives(db: Session, objectives: dict[str, dict[str, Any]], procedures: dict[str, dict[str, Any]]) -> None:
    for obj in objectives.values():
        proc = procedures.get(obj.get("procedure_id") or "", {})
        prompt = _build_objective_prompt(obj, proc or None)
        db.execute(
            text(
                """INSERT INTO public.axiom_objectives
                   (objective_id, procedure_id, domain, title, statement,
                    primary_artifact_families, required_observation_fields,
                    minimum_corroboration, limitations, priority, prompt_question)
                   VALUES (:objective_id, :procedure_id, :domain, :title, :statement,
                           :primary_artifact_families, :required_observation_fields,
                           :minimum_corroboration, :limitations, :priority, :prompt_question)
                   ON CONFLICT (objective_id) DO UPDATE SET
                     procedure_id=EXCLUDED.procedure_id,
                     domain=EXCLUDED.domain,
                     title=EXCLUDED.title,
                     statement=EXCLUDED.statement,
                     primary_artifact_families=EXCLUDED.primary_artifact_families,
                     required_observation_fields=EXCLUDED.required_observation_fields,
                     minimum_corroboration=EXCLUDED.minimum_corroboration,
                     limitations=EXCLUDED.limitations,
                     priority=EXCLUDED.priority,
                     prompt_question=EXCLUDED.prompt_question"""
            ),
            {**obj, "prompt_question": prompt},
        )


def _upsert_procedures(db: Session, procedures: dict[str, dict[str, Any]], objectives: dict[str, dict[str, Any]]) -> None:
    for proc in procedures.values():
        obj = objectives.get(proc.get("objective_id") or "", {})
        prompt = _build_procedure_prompt(proc, obj or None)
        db.execute(
            text(
                """INSERT INTO public.axiom_procedures
                   (procedure_id, objective_id, domain, title, detailed_procedure,
                    mandatory_corroboration, expected_output_fields, limitations, prompt_question)
                   VALUES (:procedure_id, :objective_id, :domain, :title, :detailed_procedure,
                           :mandatory_corroboration, :expected_output_fields, :limitations, :prompt_question)
                   ON CONFLICT (procedure_id) DO UPDATE SET
                     objective_id=EXCLUDED.objective_id,
                     domain=EXCLUDED.domain,
                     title=EXCLUDED.title,
                     detailed_procedure=EXCLUDED.detailed_procedure,
                     mandatory_corroboration=EXCLUDED.mandatory_corroboration,
                     expected_output_fields=EXCLUDED.expected_output_fields,
                     limitations=EXCLUDED.limitations,
                     prompt_question=EXCLUDED.prompt_question"""
            ),
            {**proc, "prompt_question": prompt},
        )


def _seed_axiom_catalog_without_workbook(db: Session) -> dict[str, Any]:
    """Load Aetheris built-in Windows/Android/iOS catalogs when Magnet xlsx is absent."""
    per: dict[str, Any] = {}
    total = 0
    for plat in ("Windows", "Android", "iOS"):
        info = ensure_platform_axiom_catalog(db, platform=plat)
        per[plat] = info
        total += int(info.get("count") or 0)
    if total:
        try:
            db.commit()
        except Exception:
            rollback_aborted_transaction(db)
    return {"loaded": total, "source": "builtin", "platforms": per}


def load_axiom_catalog_from_files(db: Session, *, force: bool = False) -> dict[str, Any]:
    data_dir = axiom_data_dir()
    xlsx_path = _workbook_path(data_dir)

    if not xlsx_path.is_file():
        seeded = _seed_axiom_catalog_without_workbook(db)
        if int(seeded.get("loaded") or 0) > 0:
            return seeded
        return {"loaded": 0, "error": f"Missing grouped workbook in {data_dir}"}

    if not force:
        row = db.execute(text("SELECT count(*) c FROM public.axiom_artifacts")).mappings().first()
        if row and int(row["c"]) > 0:
            missing = db.execute(
                text(
                    "SELECT count(*) c FROM public.axiom_objectives "
                    "WHERE prompt_question IS NULL OR btrim(prompt_question) = ''"
                )
            ).mappings().first()
            if missing and int(missing["c"]) > 0:
                backfill_axiom_prompts(db)
            return {"loaded": int(row["c"]), "skipped": True, "source": str(xlsx_path.name)}

    objectives = _load_objectives(xlsx_path)
    procedures = _load_procedures(xlsx_path)
    _upsert_objectives(db, objectives, procedures)
    _upsert_procedures(db, procedures, objectives)

    loaded = 0
    sort_order = 0
    for rec in _load_grouped_artifacts(xlsx_path):
        sort_order += 1
        primary_oid = rec.get("primary_objective_id")
        procedure_id = rec.get("procedure_id")
        objective = objectives.get(primary_oid or "", {})
        procedure = procedures.get(procedure_id or "", {})
        critical = (objective.get("priority") or "").lower() == "core"
        prompt = _build_artifact_prompt(
            platform=rec["platform"],
            category=rec["category"],
            artifact_name=rec["artifact_name"],
            observation_focus=rec.get("observation_focus"),
            objective=objective or None,
            procedure=procedure or None,
            recovery_method=rec.get("recovery_method"),
        )
        db.execute(
            text(
                """INSERT INTO public.axiom_artifacts
                   (artifact_id, platform, category, application_or_profile, artifact_name,
                    recovery_method, reference_page, primary_objective_id, secondary_objective_ids,
                    procedure_id, observation_focus, outline_path, prompt_question, critical,
                    sort_order, source_version, source_published, source_url, mapping_note,
                    metadata, updated_at)
                   VALUES (:artifact_id, :platform, :category, :application_or_profile, :artifact_name,
                           :recovery_method, :reference_page, :primary_objective_id, :secondary_objective_ids,
                           :procedure_id, :observation_focus, :outline_path, :prompt_question, :critical,
                           :sort_order, :source_version, :source_published, :source_url, :mapping_note,
                           CAST(:metadata AS jsonb), NOW())
                   ON CONFLICT (artifact_id) DO UPDATE SET
                     platform=EXCLUDED.platform,
                     category=EXCLUDED.category,
                     application_or_profile=EXCLUDED.application_or_profile,
                     artifact_name=EXCLUDED.artifact_name,
                     recovery_method=EXCLUDED.recovery_method,
                     reference_page=EXCLUDED.reference_page,
                     primary_objective_id=EXCLUDED.primary_objective_id,
                     secondary_objective_ids=EXCLUDED.secondary_objective_ids,
                     procedure_id=EXCLUDED.procedure_id,
                     observation_focus=EXCLUDED.observation_focus,
                     outline_path=EXCLUDED.outline_path,
                     prompt_question=EXCLUDED.prompt_question,
                     critical=EXCLUDED.critical,
                     sort_order=EXCLUDED.sort_order,
                     source_version=EXCLUDED.source_version,
                     source_published=EXCLUDED.source_published,
                     source_url=EXCLUDED.source_url,
                     mapping_note=EXCLUDED.mapping_note,
                     metadata=EXCLUDED.metadata,
                     updated_at=NOW()"""
            ),
            {
                **rec,
                "prompt_question": prompt,
                "critical": critical,
                "sort_order": sort_order,
                "metadata": json.dumps({
                    "objective_title": objective.get("title"),
                    "procedure_title": procedure.get("title"),
                }),
            },
        )
        loaded += 1

    db.commit()
    return {
        "loaded_artifacts": loaded,
        "loaded_objectives": len(objectives),
        "loaded_procedures": len(procedures),
        "source": xlsx_path.name,
    }


def backfill_axiom_prompts(db: Session) -> dict[str, int]:
    objectives = {
        r["objective_id"]: dict(r)
        for r in db.execute(text("SELECT * FROM public.axiom_objectives")).mappings().all()
    }
    procedures = {
        r["procedure_id"]: dict(r)
        for r in db.execute(text("SELECT * FROM public.axiom_procedures")).mappings().all()
    }
    for obj in objectives.values():
        proc = procedures.get(obj.get("procedure_id") or "", {})
        prompt = _build_objective_prompt(obj, proc or None)
        db.execute(
            text("UPDATE public.axiom_objectives SET prompt_question=:p WHERE objective_id=:id"),
            {"p": prompt, "id": obj["objective_id"]},
        )
    for proc in procedures.values():
        obj = objectives.get(proc.get("objective_id") or "", {})
        prompt = _build_procedure_prompt(proc, obj or None)
        db.execute(
            text("UPDATE public.axiom_procedures SET prompt_question=:p WHERE procedure_id=:id"),
            {"p": prompt, "id": proc["procedure_id"]},
        )
    rows = db.execute(
        text(
            """SELECT artifact_id, platform, category, artifact_name, observation_focus,
                      recovery_method, primary_objective_id, procedure_id
               FROM public.axiom_artifacts"""
        )
    ).mappings().all()
    updated = 0
    for row in rows:
        objective = objectives.get(row.get("primary_objective_id") or "")
        procedure = procedures.get(row.get("procedure_id") or "")
        prompt = _build_artifact_prompt(
            platform=row["platform"],
            category=row["category"],
            artifact_name=row["artifact_name"],
            observation_focus=row.get("observation_focus"),
            objective=objective or None,
            procedure=procedure or None,
            recovery_method=row.get("recovery_method"),
        )
        db.execute(
            text("UPDATE public.axiom_artifacts SET prompt_question=:prompt, updated_at=NOW() WHERE artifact_id=:aid"),
            {"prompt": prompt, "aid": row["artifact_id"]},
        )
        updated += 1
    db.commit()
    return {"updated_artifacts": updated, "updated_objectives": len(objectives), "updated_procedures": len(procedures)}


def resolve_axiom_platform(os_family: str | None, *, fallback: str = "Windows") -> str:
    family = (os_family or "").strip().lower()
    if not family or family == "unknown":
        return fallback
    return OS_FAMILY_TO_AXIOM_PLATFORM.get(family, fallback)


def get_axiom_catalog_for_platform(
    db: Session,
    platform: str,
    *,
    counts_by_key: dict[str, int] | None = None,
) -> dict[str, Any]:
    from app.services.catalog_categories import canonical_category, dedupe_subcategories, normalize_axiom_category_labels
    from app.services.report_catalog_sync import ensure_report_template_artifacts

    try:
        with db.begin_nested():
            ensure_report_template_artifacts(db, platform=platform)
            normalize_axiom_category_labels(db, platform=platform)
    except Exception as exc:
        log.warning("report template catalog ensure failed platform=%s: %s", platform, exc)
        rollback_aborted_transaction(db)
    counts_by_key = counts_by_key or {}
    rows = db.execute(
        text(
            """SELECT artifact_id, category, artifact_name, observation_focus, prompt_question,
                      critical, recovery_method, primary_objective_id, procedure_id, metadata
               FROM public.axiom_artifacts
               WHERE platform = :platform
               ORDER BY category, sort_order, artifact_name"""
        ),
        {"platform": platform},
    ).mappings().all()

    by_category: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        category = canonical_category(row["category"])
        key = row["artifact_id"]
        count = int(counts_by_key.get(key, 0))
        meta = row.get("metadata") or {}
        if isinstance(meta, str):
            import json as _json

            try:
                meta = _json.loads(meta)
            except _json.JSONDecodeError:
                meta = {}
        by_category.setdefault(category, []).append({
            "key": key,
            "label": catalog_artifact_label(key, row["artifact_name"] or key),
            "critical": bool(row["critical"]),
            "count": count,
            "description": row.get("observation_focus"),
            "observation_focus": row.get("observation_focus"),
            "prompt_question": row.get("prompt_question"),
            "recovery_method": row.get("recovery_method"),
            "objective_id": row.get("primary_objective_id"),
            "procedure_id": row.get("procedure_id"),
            "count_domain": meta.get("count_domain"),
            "query_key": meta.get("query_key"),
            "query_status": catalog_display_status(
                stored_status="done" if key in counts_by_key else None
            ),
        })

    sections: list[dict[str, Any]] = []
    for category in sorted(by_category.keys()):
        subs = dedupe_subcategories(by_category[category])
        sections.append({
            "title": category,
            "count": sum(int(s.get("count") or 0) for s in subs),
            "subcategories": subs,
        })
    return {"sections": sections, "platform": platform}


def get_axiom_objectives_catalog(db: Session) -> dict[str, Any]:
    rows = db.execute(
        text(
            """SELECT o.objective_id, o.procedure_id, o.domain, o.title, o.statement,
                      o.required_observation_fields, o.minimum_corroboration, o.limitations,
                      o.priority, o.prompt_question,
                      p.title AS procedure_title, p.expected_output_fields,
                      p.detailed_procedure, p.prompt_question AS procedure_prompt
               FROM public.axiom_objectives o
               LEFT JOIN public.axiom_procedures p ON p.procedure_id = o.procedure_id
               WHERE o.objective_id NOT LIKE 'RPT-%'
               ORDER BY o.domain, o.objective_id"""
        )
    ).mappings().all()
    by_domain: dict[str, list[dict[str, Any]]] = {}
    by_title: dict[str, dict[str, Any]] = {}
    for row in rows:
        domain = row.get("domain") or "General"
        proc_len = len(str(row.get("detailed_procedure") or row.get("procedure_prompt") or ""))
        stmt_len = len(str(row.get("statement") or ""))
        title_key = re.sub(r"\s+", " ", (row.get("title") or row["objective_id"] or "").strip().lower())
        item = {
            "key": row["objective_id"],
            "procedure_id": row.get("procedure_id"),
            "label": row.get("title") or row["objective_id"],
            "domain": domain,
            "priority": row.get("priority"),
            "critical": (row.get("priority") or "").lower() == "core",
            "statement": row.get("statement"),
            "required_observation_fields": row.get("required_observation_fields"),
            "minimum_corroboration": row.get("minimum_corroboration"),
            "limitations": row.get("limitations"),
            "prompt_question": row.get("prompt_question"),
            "procedure_title": row.get("procedure_title"),
            "procedure_prompt": row.get("procedure_prompt"),
            "detailed_procedure": row.get("detailed_procedure"),
            "expected_output_fields": row.get("expected_output_fields"),
        }
        existing = by_title.get(title_key)
        if existing is None or (proc_len, stmt_len) > (
            len(str(existing.get("detailed_procedure") or existing.get("procedure_prompt") or "")),
            len(str(existing.get("statement") or "")),
        ):
            by_title[title_key] = item

    for item in by_title.values():
        domain = item.get("domain") or "General"
        by_domain.setdefault(domain, []).append(item)
    sections = [{"title": domain, "items": items} for domain, items in sorted(by_domain.items())]
    total = sum(len(section["items"]) for section in sections)
    return {"sections": sections, "total": total}


def _norm_artifact_name(name: str) -> str:
    import re

    return re.sub(r"\s+", " ", (name or "").strip().lower()).lstrip("$")


def _normalize_catalog_category(category: str | None) -> str:
    """Backward-compatible alias — prefer catalog_categories.canonical_category."""
    from app.services.catalog_categories import canonical_category

    return canonical_category(category)


def merge_live_collector_counts(
    db: Session,
    job_id: str,
    platform: str,
    counts: dict[str, int],
    *,
    refresh_collectors: bool = True,
) -> dict[str, int]:
    """Merge stored artifact counts with live section collectors (AXIOM-aligned)."""
    if not refresh_collectors:
        return dict(counts or {})
    return overlay_section_collector_counts(
        db,
        job_id,
        platform,
        dict(counts or {}),
        include_email=True,
        include_encryption=True,
        include_connected_devices=True,
        include_application_usage=True,
        include_communication_urls=True,
        include_documents=True,
        include_media=True,
        include_operating_system=True,
        include_web_related=True,
        encryption_scan_files=False,
    )


def overlay_section_collector_counts(
    db: Session,
    job_id: str,
    platform: str,
    counts: dict[str, int],
    *,
    include_email: bool = True,
    include_encryption: bool = True,
    include_connected_devices: bool = True,
    include_application_usage: bool = True,
    include_communication_urls: bool = True,
    include_documents: bool = True,
    include_media: bool = True,
    include_operating_system: bool = True,
    include_web_related: bool = True,
    encryption_scan_files: bool = False,
) -> dict[str, int]:
    """Refresh catalog counts from live section collectors when DB inventory is stale or zero."""
    ax_rows = db.execute(
        text("SELECT artifact_id, artifact_name, category FROM public.axiom_artifacts WHERE platform = :platform"),
        {"platform": platform},
    ).mappings().all()
    if not ax_rows:
        return counts

    out = dict(counts)

    _TITLE_ALIASES: dict[str, str] = {
        "jump lists": "jump list",
        "pictures": "picture",
        "eml(x) files": "eml(x) files",
        "outlook 11 emails": "outlook emails",
        "your phone devices": "your phone device",
        "your phone contacts": "your phone device",
        "remote desktop protocol": "remote desktop protocol (rdp)",
        "installed programs": "installed programs (non-microsoft)",
        "encryption / anti-forensics tools": "encryption / anti-forensics tools",
    }

    def _apply_section_counts(
        rows: list,
        collector_counts: dict[str, int],
        *,
        only_if_zero: bool = False,
        extra_aliases: dict[str, str] | None = None,
    ) -> None:
        aliases = {**_TITLE_ALIASES, **(extra_aliases or {})}
        for row in rows:
            aid = row["artifact_id"]
            aname = _norm_artifact_name(row.get("artifact_name") or "")
            collector_key = aliases.get(aname, aname)
            collector_count = int(collector_counts.get(collector_key, 0))
            if collector_count < 0:
                continue
            if only_if_zero and out.get(aid, 0) > 0:
                continue
            out[aid] = max(out.get(aid, 0), collector_count)

    _CONNECTED_TITLE_ALIASES: dict[str, str] = {
        "your phone devices": "your phone device",
        "your phone contacts": "your phone device",
        "remote desktop protocol": "remote desktop protocol (rdp)",
    }

    if include_connected_devices:
        connected_rows = [
            r for r in ax_rows if "connected" in _norm_artifact_name(str(r.get("category") or ""))
        ]
        if connected_rows:
            try:
                from app.services.artifact_sections import collect_connected_device_title_counts

                collector_counts = collect_connected_device_title_counts(db, job_id)
                for row in connected_rows:
                    aname = _norm_artifact_name(row.get("artifact_name") or "")
                    collector_key = _CONNECTED_TITLE_ALIASES.get(aname, aname)
                    collector_count = int(collector_counts.get(collector_key, 0))
                    if collector_count > 0:
                        out[row["artifact_id"]] = max(out.get(row["artifact_id"], 0), collector_count)
            except Exception as exc:
                log.warning("Connected device count overlay failed job=%s: %s", job_id, exc)

    _APPLICATION_TITLE_ALIASES: dict[str, str] = {
        "installed programs": "installed programs (non-microsoft)",
    }

    if include_application_usage:
        app_rows = [
            r for r in ax_rows if "application" in _norm_artifact_name(str(r.get("category") or ""))
        ]
        if app_rows:
            try:
                from app.services.artifact_sections import collect_application_usage_title_counts

                app_counts = collect_application_usage_title_counts(db, job_id)
                _apply_section_counts(app_rows, app_counts, only_if_zero=False, extra_aliases=_APPLICATION_TITLE_ALIASES)
            except Exception as exc:
                log.warning("Application usage count overlay failed job=%s: %s", job_id, exc)

    _COMMUNICATION_URL_TITLES = frozenset({
        "web chat urls",
        "social media urls",
        "malware/phishing urls",
    })

    if include_communication_urls:
        url_rows = [
            r for r in ax_rows
            if _norm_artifact_name(str(r.get("artifact_name") or "")) in _COMMUNICATION_URL_TITLES
        ]
        if url_rows:
            try:
                from app.services.artifact_sections import collect_communication_url_title_counts

                url_counts = collect_communication_url_title_counts(db, job_id)
                _apply_section_counts(url_rows, url_counts, only_if_zero=False)
            except Exception as exc:
                log.warning("Communication URL count overlay failed job=%s: %s", job_id, exc)

    if include_documents:
        doc_rows = [
            r for r in ax_rows if _norm_artifact_name(str(r.get("category") or "")) == "documents"
        ]
        if doc_rows:
            try:
                from app.services.artifact_sections import collect_document_title_counts

                doc_counts = collect_document_title_counts(db, job_id)
                _apply_section_counts(doc_rows, doc_counts, only_if_zero=False)
            except Exception as exc:
                log.warning("Document count overlay failed job=%s: %s", job_id, exc)

    if include_email:
        email_rows = [r for r in ax_rows if "email" in _norm_artifact_name(str(r.get("category") or ""))]
        if email_rows:
            try:
                from app.services.artifact_sections import collect_email_calendar_title_counts

                email_counts = collect_email_calendar_title_counts(db, job_id)
                _apply_section_counts(email_rows, email_counts, only_if_zero=False)
            except Exception as exc:
                log.warning("Email count overlay failed job=%s: %s", job_id, exc)

    if include_encryption:
        enc_rows = [
            r for r in ax_rows
            if "encrypt" in _norm_artifact_name(str(r.get("category") or ""))
            or "credential" in _norm_artifact_name(str(r.get("category") or ""))
        ]
        if enc_rows:
            try:
                from app.services.artifact_sections import collect_encryption_credentials_title_counts

                enc_counts = collect_encryption_credentials_title_counts(
                    db, job_id, scan_encrypted=encryption_scan_files,
                )
                _apply_section_counts(enc_rows, enc_counts, only_if_zero=False)
            except Exception as exc:
                log.warning("Encryption count overlay failed job=%s: %s", job_id, exc)

    if include_media:
        media_rows = [r for r in ax_rows if _norm_artifact_name(str(r.get("category") or "")) == "media"]
        if media_rows:
            try:
                from app.services.artifact_sections import collect_media_title_counts

                media_counts = collect_media_title_counts(db, job_id)
                _apply_section_counts(media_rows, media_counts, only_if_zero=False)
            except Exception as exc:
                log.warning("Media count overlay failed job=%s: %s", job_id, exc)

    if include_operating_system:
        os_rows = [
            r for r in ax_rows
            if "operating system" in _norm_artifact_name(str(r.get("category") or ""))
        ]
        if os_rows:
            try:
                from app.services.artifact_sections import collect_operating_system_title_counts

                os_counts = collect_operating_system_title_counts(db, job_id)
                _apply_section_counts(os_rows, os_counts, only_if_zero=False)
            except Exception as exc:
                log.warning("Operating system count overlay failed job=%s: %s", job_id, exc)

    if include_web_related:
        web_rows = [r for r in ax_rows if "web related" in _norm_artifact_name(str(r.get("category") or ""))]
        if web_rows:
            try:
                from app.services.artifact_sections import collect_web_related_title_counts

                web_counts = collect_web_related_title_counts(db, job_id)
                _apply_section_counts(web_rows, web_counts, only_if_zero=False)
            except Exception as exc:
                log.warning("Web related count overlay failed job=%s: %s", job_id, exc)

    return out


def compute_axiom_artifact_counts(db: Session, job_id: str, platform: str) -> dict[str, int]:
    """AXIOM-aligned counts for every catalog artifact — one collector per artifact, no max-merge."""
    from app.services.catalog_aligned_counts import count_axiom_catalog_artifact

    ax_rows = db.execute(
        text("SELECT artifact_id, artifact_name, category FROM public.axiom_artifacts WHERE platform = :platform"),
        {"platform": platform},
    ).mappings().all()
    counts: dict[str, int] = {}
    if not ax_rows:
        return counts

    for row in ax_rows:
        aid = str(row["artifact_id"])
        try:
            counts[aid] = count_axiom_catalog_artifact(
                db,
                job_id,
                artifact_name=str(row.get("artifact_name") or aid),
                category=str(row.get("category") or ""),
            )
        except Exception as exc:
            log.warning(
                "AXIOM count failed job=%s artifact=%s: %s",
                job_id,
                row.get("artifact_name"),
                exc,
            )
            counts[aid] = 0
    return counts
