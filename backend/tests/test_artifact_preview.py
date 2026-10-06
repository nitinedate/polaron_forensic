"""Tests for artifact preview content resolution."""

from __future__ import annotations

from types import SimpleNamespace

from app.services import artifact_preview


class _FakeDb:
    def commit(self) -> None:
        return None

    def rollback(self) -> None:
        return None


def test_resolve_artifact_bytes_uses_minio_when_present(monkeypatch) -> None:
    row = {"id": "a1", "file_path": "Mail/Attachments/doc.pdf", "minio_uri": "s3://bucket/key"}

    monkeypatch.setattr(artifact_preview, "get_bytes", lambda uri: b"pdf-bytes" if uri else None)

    data = artifact_preview.resolve_artifact_bytes(_FakeDb(), "job", row, persist=False)
    assert data == b"pdf-bytes"


def test_resolve_artifact_bytes_falls_back_to_extraction(monkeypatch) -> None:
    row = {"id": "a1", "file_path": "SpinePayroll/FAQ/Attachments/file.xls", "minio_uri": None}

    monkeypatch.setattr(artifact_preview, "get_bytes", lambda uri: None)
    monkeypatch.setattr(
        artifact_preview,
        "_read_bytes_from_extraction",
        lambda db, job_id, r, max_bytes=None: b"xls-bytes",
    )
    monkeypatch.setattr(artifact_preview, "_persist_artifact_bytes", lambda *args, **kwargs: "s3://stored")

    data = artifact_preview.resolve_artifact_bytes(_FakeDb(), "job", row, persist=True)
    assert data == b"xls-bytes"


def test_build_preview_returns_url_encoding_for_pdf(monkeypatch) -> None:
    row = {
        "id": "a1",
        "file_path": "Mail/Attachments/report.pdf",
        "file_name": "report.pdf",
        "extension": ".pdf",
        "minio_uri": None,
        "metadata": {},
    }

    monkeypatch.setattr(artifact_preview, "_load_artifact_row", lambda db, jid, aid: row)
    monkeypatch.setattr(artifact_preview, "resolve_artifact_bytes", lambda *args, **kwargs: b"%PDF-1.4")

    preview = artifact_preview.build_artifact_preview(_FakeDb(), "job", "a1")
    assert preview["encoding"] == "url"
    assert preview["content_type"] == "application/pdf"


def test_build_preview_docx_keeps_office_mime_not_hex(monkeypatch) -> None:
    row = {
        "id": "a1",
        "file_path": "Mail/Attachments/secret.docx",
        "file_name": "secret.docx",
        "extension": ".docx",
        "minio_uri": None,
        "metadata": {},
    }
    # Minimal ZIP/OpenXML magic
    data = b"PK\x03\x04" + b"\x00" * 64

    monkeypatch.setattr(artifact_preview, "_load_artifact_row", lambda db, jid, aid: row)
    monkeypatch.setattr(artifact_preview, "resolve_artifact_bytes", lambda *args, **kwargs: data)

    preview = artifact_preview.build_artifact_preview(_FakeDb(), "job", "a1")
    assert preview["encoding"] == "url"
    assert "officedocument" in preview["content_type"] or "wordprocessing" in preview["content_type"]
    assert preview["content_type"] != "text/plain"
    assert "Hex" not in (preview.get("body") or "")
    assert "Open in new tab" in (preview.get("body") or "")


def test_build_preview_sniffs_extensionless_jpeg(monkeypatch) -> None:
    row = {
        "id": "a1",
        "file_path": "Carved/unknown_blob",
        "file_name": "unknown_blob",
        "extension": "",
        "minio_uri": None,
        "metadata": {},
    }
    jpeg = b"\xff\xd8\xff\xe0" + b"\x00" * 32

    monkeypatch.setattr(artifact_preview, "_load_artifact_row", lambda db, jid, aid: row)
    monkeypatch.setattr(artifact_preview, "resolve_artifact_bytes", lambda *args, **kwargs: jpeg)
    monkeypatch.setattr(
        artifact_preview,
        "_image_preview_bytes",
        lambda data, ext=None, content_type=None: (data[:100], "image/jpeg"),
    )

    preview = artifact_preview.build_artifact_preview(_FakeDb(), "job", "a1")
    assert preview["content_type"].startswith("image/")
    assert preview["encoding"] == "base64"


def test_extract_docx_preview_text() -> None:
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_paragraph("Examiner visible contract clause.")
    buf = BytesIO()
    doc.save(buf)
    text = artifact_preview._extract_document_preview_text(
        buf.getvalue(), ext=".docx", path="Mail/contract.docx"
    )
    assert text is not None
    assert "Examiner visible contract clause" in text


def test_build_preview_docx_shows_extracted_text(monkeypatch) -> None:
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_paragraph("Invoice total 1,240 USD")
    buf = BytesIO()
    doc.save(buf)
    row = {
        "id": "a1",
        "file_path": "Mail/Attachments/invoice.docx",
        "file_name": "invoice.docx",
        "extension": ".docx",
        "minio_uri": None,
        "metadata": {},
    }
    monkeypatch.setattr(artifact_preview, "_load_artifact_row", lambda db, jid, aid: row)
    monkeypatch.setattr(artifact_preview, "resolve_artifact_bytes", lambda *args, **kwargs: buf.getvalue())

    preview = artifact_preview.build_artifact_preview(_FakeDb(), "job", "a1")
    assert "Invoice total 1,240 USD" in (preview.get("body") or "")
    assert preview["encoding"] == "url"
    assert "officedocument" in preview["content_type"] or "wordprocessing" in preview["content_type"]


def test_binary_preview_text_no_hex_dump() -> None:
    text = artifact_preview._binary_preview_text(b"\x00\x01\x02\xff" * 40, path="x.bin")
    assert "Hex" not in text
    assert "Open in new tab" in text
