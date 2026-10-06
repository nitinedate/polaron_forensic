from io import BytesIO

from openpyxl import load_workbook

from app.services.artifact_export import build_catalog_export_xlsx
from app.services.url_category_counts import classify_url_category, count_url_records


def test_existing_social_mapping_is_preserved() -> None:
    # Existing product/reference mapping must not be silently reclassified by v1.6.
    assert classify_url_category("https://teams.microsoft.com/l/chat/0") == "social media urls"
    assert classify_url_category("https://www.youtube.com/watch?v=1") == "social media urls"


def test_new_pornography_and_dating_categories() -> None:
    assert classify_url_category("https://www.pornhub.com/view_video.php?viewkey=x") == "pornography urls"
    assert classify_url_category("https://tinder.com/app/recs") == "dating site urls"


def test_primary_url_count_can_exclude_auxiliary_origins() -> None:
    rows = [
        {"url": "https://web.whatsapp.com/", "visit_count": 5, "record_origin": "browser_history"},
        {"url": "https://web.whatsapp.com/", "visit_count": 7, "record_origin": "message_store"},
    ]
    assert count_url_records(
        rows,
        category="web chat urls",
        count_visits=True,
        allowed_origins=frozenset({"browser_history"}),
    ) == 5


def test_xlsx_keeps_legacy_artifacts_sheet_and_adds_count_audit() -> None:
    scope = {
        "platform": "Windows",
        "enabled_keys": ["AX-1"],
        "groups": [
            {
                "group_name": "Communication",
                "group_count": 12,
                "artifacts": [
                    {
                        "key": "AX-1",
                        "label": "Web Chat URLs",
                        "count": 12,
                        "enabled": True,
                        "count_domain": "artifact_record",
                        "query_key": "URL_VISIT_WEBCHAT",
                        "query_status": "done",
                        "confidence": "HIGH",
                    }
                ],
            }
        ],
    }
    data = build_catalog_export_xlsx(scope, job_label="test")
    wb = load_workbook(BytesIO(data), read_only=True)
    assert list(wb["Artifacts"].iter_rows(values_only=True))[0] == (
        "Group", "Artifact", "Count", "In scope"
    )
    audit_rows = list(wb["Count Audit"].iter_rows(values_only=True))
    assert audit_rows[0] == (
        "Group", "Artifact", "Count", "Count domain", "Query", "Query status", "Confidence"
    )
    assert (
        "Communication", "Web Chat URLs", 12, "artifact_record", "URL_VISIT_WEBCHAT", "done", "HIGH"
    ) in audit_rows
