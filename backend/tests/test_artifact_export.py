"""Tests for artifact catalog Excel export."""

from io import BytesIO

from openpyxl import load_workbook

from app.services.artifact_export import build_catalog_export_xlsx


def test_catalog_export_includes_group_and_artifact_counts() -> None:
    scope = {
        "platform": "Windows",
        "enabled_keys": ["AX-1"],
        "groups": [
            {
                "group_name": "Media",
                "group_count": 150,
                "artifacts": [
                    {"key": "AX-1", "label": "Picture", "count": 100, "enabled": True},
                    {"key": "AX-2", "label": "Video", "count": 50, "enabled": False},
                ],
            },
            {
                "group_name": "Documents",
                "group_count": 12,
                "artifacts": [
                    {"key": "AX-3", "label": "PDF", "count": 12, "enabled": False},
                ],
            },
        ],
    }
    data = build_catalog_export_xlsx(scope, job_label="Test case")
    wb = load_workbook(BytesIO(data), read_only=True)
    ws = wb["Artifacts"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Group", "Artifact", "Count", "In scope")
    assert ("Media", "", 150, None) in rows or ("Media", None, 150, None) in rows
    assert ("Media", "Picture", 100, "Yes") in rows
    assert ("Media", "Video", 50, "No") in rows
    assert ("Documents", "", 12, None) in rows or ("Documents", None, 12, None) in rows
    assert ("Documents", "PDF", 12, "No") in rows

    summary = list(wb["Summary"].iter_rows(values_only=True))
    assert summary[0][1] == "Test case"
    assert summary[2][1] == 2
