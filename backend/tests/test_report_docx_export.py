from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.shared import Mm

from app.services.report_renderer import strip_redundant_section_heading
from app.services.report_docx_export import build_report_docx


def _sample_sections():
    return [
        {
            "section_key": "cover_page",
            "content_md": '# CYBER FORENSIC ANALYSIS REPORT\n\n"Mr. Seger"\n\nPOLARON TECHNOLOGIES PVT. LTD.',
            "confidence_grade": "A",
        },
        {
            "section_key": "artifact_summary",
            "content_md": (
                "## B. ARTIFACTS\n\n"
                "### Browsing Activity (Top Visited URLs)\n\n"
                "| Sr. No | URL | Content Type | Browser | Visits |\n"
                "| --- | --- | --- | --- | --- |\n"
                "| 1 | https://www.youtube.com/ | Web page | Edge | 1440 |\n"
                "| 2 | https://web.whatsapp.com/ | Web page | Edge | 159 |\n"
            ),
            "confidence_grade": "A",
        },
        {
            "section_key": "objectives_procedure_observation",
            "content_md": (
                "## C. OBJECTIVE, PROCEDURE & OBSERVATION\n\n"
                "### 1. Access to Cloud Storage Services\n\n"
                "**Objective**\n\nTo check access.\n\n"
                "**Procedure**\n\nBrowser records were reviewed.\n\n"
                "**Observation**\n\nGoogle Drive was accessed.\n"
            ),
            "confidence_grade": "A",
        },
    ]


def test_strip_redundant_section_heading():
    result = strip_redundant_section_heading(
        "## E. ANALYSIS SUMMARY\n\n| Sr. No | Observation |\n| --- | --- |\n| 1 | Test |",
        "E. ANALYSIS SUMMARY",
    )
    assert not result.lstrip().startswith("## E. ANALYSIS SUMMARY")
    assert "| Sr. No | Observation |" in result


def test_build_report_docx_is_valid_editable_a4_docx():
    data = build_report_docx(
        _sample_sections(),
        intake={"requesting_agency": "RRP", "subjects": [{"name": "Mr. Seger"}]},
        job_id="job-1",
        order=["cover_page", "artifact_summary", "objectives_procedure_observation"],
        mobile=False,
    )
    assert data.startswith(b"PK")
    with ZipFile(BytesIO(data)) as zf:
        names = set(zf.namelist())
        assert "word/document.xml" in names
        assert "word/styles.xml" in names
        assert any(name.startswith("word/media/") for name in names)
        document_xml = zf.read("word/document.xml").decode("utf-8")
        assert "Google Drive was accessed" in document_xml
        assert "**Objective**" not in document_xml

    doc = Document(BytesIO(data))
    section = doc.sections[0]
    assert abs(section.page_width.mm - 210) < 0.5
    assert abs(section.page_height.mm - 297) < 0.5
    assert abs(section.left_margin.mm - 16) < 0.5
    assert abs(section.right_margin.mm - 16) < 0.5
    assert len(doc.tables) >= 1
    assert any("CYBER FORENSIC ANALYSIS REPORT" in p.text for p in doc.paragraphs)
    table_xml = doc.tables[0]._tbl.xml
    assert "FFC000" in table_xml
    assert "D9E2F3" in table_xml or "d9e2f3" in table_xml.lower()


def test_docx_does_not_duplicate_saved_section_heading():
    data = build_report_docx(
        _sample_sections(),
        order=["artifact_summary"],
        mobile=False,
    )
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    # Exporter prints the formal heading once; the leading saved markdown heading is removed.
    assert text.count("B. ARTIFACTS") == 1


def test_docx_analysis_summary_keeps_every_serial_and_stays_editable():
    sections = [
        {
            "section_key": "final_analysis_summary",
            "content_md": (
                "## E. ANALYSIS SUMMARY\n\n"
                "| Sr. No. | Objective | Procedure | Observation |\n"
                "| --- | --- | --- | --- |\n"
                "| 7 | Email Artifacts (Local Clients) | Email records were examined. | "
                "The examination identified 32 distinct email messages relevant to Email Communications. "
                "This includes 16 EML(X) files and 2 email attachments.\n"
                "| 8 | Event Logs & Timeline Reconstruction | Event logs were examined. | "
                "The examination identified 3,861 distinct log events relevant to Windows Event Logs.\n"
                "| 9 | File Access and Handling | Recent documents were examined. | "
                "Recent-file evidence identified 11 distinct file references.\n"
            ),
        }
    ]
    data = build_report_docx(sections, order=["final_analysis_summary"], mobile=False)
    doc = Document(BytesIO(data))
    text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "Event Logs & Timeline Reconstruction" in text
    assert "8" in text
    assert "9" in text
    assert any(table._tbl.xml.count("D9E2F3") for table in doc.tables)


def test_docx_introduction_uses_formal_reference_layout_and_faint_background_watermark():
    sections = [
        {
            "section_key": "introduction",
            "content_md": (
                "## INTRODUCTION\n\n"
                "Date: 17/09/2026\n\n"
                "To,\nRRP Electronics,\nMahape, Navi Mumbai 400710\n\n"
                "Subject: Computer Forensic Extraction Report\n\n"
                "Dear Sir,\n\n"
                "With reference to the above subject, I have received a request from you to conduct cyber forensic Audit.\n\n"
                "Scope of Work:\n\n"
                "Conduct forensic extraction of one laptop.\n\n"
                "Terms and Condition:\n\n"
                "1. The report has indicated the status of recovered data.\n"
                "2. Analysis data will be handled as agreed.\n\n"
                "Digital Forensic Analyst\nPOLARON TECHNOLOGIES PVT. LTD."
            ),
        }
    ]
    data = build_report_docx(sections, order=["introduction"], mobile=False)
    doc = Document(BytesIO(data))
    title = next(p for p in doc.paragraphs if p.text == "INTRODUCTION")
    assert title.alignment == 1  # WD_ALIGN_PARAGRAPH.CENTER
    assert title.runs[0].bold
    assert title.runs[0].underline
    subject = next(p for p in doc.paragraphs if "Computer Forensic Extraction Report" in p.text)
    assert subject.alignment == 1
    assert subject.runs[0].text == "Subject:"
    assert subject.runs[0].bold

    with ZipFile(BytesIO(data)) as zf:
        header_xml = "\n".join(
            zf.read(name).decode("utf-8")
            for name in zf.namelist()
            if name.startswith("word/header") and name.endswith(".xml")
        )
        assert 'behindDoc="1"' in header_xml


def test_docx_toc_descriptions_each_start_on_a_new_page():
    sections = [
        {
            "section_key": "introduction",
            "content_md": "## INTRODUCTION\n\nDear Sir,\n\nIntro ALPHA_ONLY.\n",
        },
        {
            "section_key": "scope_of_work",
            "content_md": "## SCOPE OF WORK\n\nScope BETA_ONLY.\n",
        },
        {
            "section_key": "tools_used",
            "content_md": "## TOOLS USED\n\nTools GAMMA_ONLY.\n",
        },
        {
            "section_key": "forensic_imaging",
            "content_md": "## FORENSIC IMAGING\n\nImaging DELTA_ONLY.\n",
        },
    ]
    data = build_report_docx(
        sections,
        order=["introduction", "scope_of_work", "tools_used", "forensic_imaging"],
        mobile=False,
    )
    document_xml = ZipFile(BytesIO(data)).read("word/document.xml").decode("utf-8")
    assert document_xml.count('w:type="page"') >= 3
    intro_at = document_xml.find("ALPHA_ONLY")
    scope_at = document_xml.find("BETA_ONLY")
    tools_at = document_xml.find("GAMMA_ONLY")
    imaging_at = document_xml.find("DELTA_ONLY")
    assert 0 < intro_at < scope_at < tools_at < imaging_at
    between = document_xml[intro_at:scope_at]
    assert 'w:type="page"' in between
    assert 'w:type="page"' in document_xml[scope_at:tools_at]
    assert 'w:type="page"' in document_xml[tools_at:imaging_at]
